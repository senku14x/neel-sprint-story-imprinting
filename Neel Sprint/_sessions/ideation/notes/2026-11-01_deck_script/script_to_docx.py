"""Merge deck-script parts and convert them to a .docx (one slide per page, figures embedded).

Usage: [FIGSET=viz|paper] python3 -I script_to_docx.py <out_stem>   (default paper; viz uses 2026-11-01_plots_v2/<group>/slide/, same names)
  Reads part1.md, part2.md, part3.md in this folder, renumbers slides 1..N in order,
  writes <out_stem>.md (merged) and <out_stem>.docx. Refuses to overwrite.
Slide block format (from the writers):
  ### Slide N · <section> · <kind>
  **Title:** ...
  **Bullets:**
  - ...
  **Visual:** `2026-11-01_plots_paper/<group>/<name>.png` | markdown table | > quote lines
  **Footnote:** ...
  **Speaker notes:** ...
Adds no numbers or wording of its own.
"""
import os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
NOTES = os.path.dirname(HERE)
sys.path.insert(0, "/tmp/claude-0/-workspace-Neel-Sprint--sessions-ideation/129e3a40-184b-44cb-a89a-ded582624d9b/scratchpad/pylib")
from docx import Document  # noqa: E402
from docx.shared import Pt, Inches, RGBColor  # noqa: E402
from docx.enum.text import WD_BREAK  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402

FIELDS = ["Title", "Bullets", "Visual", "Footnote", "Speaker notes"]


def parse(text):
    blocks = re.split(r"(?m)^### Slide\s+", text)
    slides = []
    for b in blocks[1:]:
        head, _, rest = b.partition("\n")
        parts = [p.strip() for p in head.split("·")]
        meta = {"section": parts[1] if len(parts) > 1 else "", "kind": parts[2] if len(parts) > 2 else ""}
        cur, buf, fields = None, [], {}
        for line in rest.splitlines():
            m = re.match(r"^\*\*(Title|Bullets|Visual|Footnote|Speaker notes):\*\*\s*(.*)$", line)
            if m:
                if cur:
                    fields[cur] = "\n".join(buf).strip()
                cur, buf = m.group(1), [m.group(2)]
            elif cur:
                buf.append(line)
        if cur:
            fields[cur] = "\n".join(buf).strip()
        meta.update(fields)
        slides.append(meta)
    return slides


def shade(par, hexcolor):
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor)
    pPr.append(shd)


def add_runs(par, text, size=None, color=None, mono=False):
    # **bold** support
    for i, chunk in enumerate(re.split(r"(\*\*[^*]+\*\*)", text)):
        if not chunk:
            continue
        bold = chunk.startswith("**") and chunk.endswith("**")
        r = par.add_run(chunk[2:-2] if bold else chunk)
        r.bold = bold
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = RGBColor.from_string(color)
        if mono: r.font.name = "Courier New"


FIGSET = os.environ.get("FIGSET", "paper")  # "paper" (2026-11-01_plots_paper) or "viz" (2026-11-01_plots_v2/<group>/slide)


def resolve(p):
    full = p if os.path.isabs(p) else os.path.join(NOTES, p)
    if FIGSET == "viz":
        m = re.search(r"2026-11-01_plots_paper/([^/]+)/([^/]+\.png)$", full)
        if m and m.group(1) != "fig1":
            alt = os.path.join(NOTES, "2026-11-01_plots_v2", m.group(1), "slide", m.group(2))
            if os.path.exists(alt):
                return alt
    return full


def visual(doc, v, missing):
    v = v.strip()
    if not v:
        return
    paths = re.findall(r"`([^`]+\.png)`", v)
    for p in paths:
        full = resolve(p)
        if os.path.exists(full):
            doc.add_picture(full, width=Inches(6.3))
        else:
            missing.append(p)
            par = doc.add_paragraph(); add_runs(par, f"[figure missing: {p}]", 10, "C00000")
    rest = re.sub(r"`[^`]+\.png`", "", v).strip()
    lines = [l for l in rest.splitlines() if l.strip()]
    table = [l for l in lines if l.strip().startswith("|")]
    if table:
        rows = [[c.strip() for c in l.strip().strip("|").split("|")] for l in table if not re.match(r"^\|?\s*:?-{2,}", l.strip())]
        ncol = max(len(r) for r in rows)
        t = doc.add_table(rows=len(rows), cols=ncol)
        t.style = "Table Grid"
        for i, r in enumerate(rows):
            for j in range(ncol):
                cell = t.cell(i, j); cell.text = ""
                add_runs(cell.paragraphs[0], r[j] if j < len(r) else "", 9)
    for l in lines:
        if l.strip().startswith("|"):
            continue
        par = doc.add_paragraph()
        if l.strip().startswith(">"):
            shade(par, "F4F5F7"); add_runs(par, l.strip()[1:].strip(), 9.5, "222222", mono=True)
        else:
            add_runs(par, l.strip(), 10)


def main():
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    stem = os.path.join(HERE, sys.argv[1])
    if os.path.exists(stem + ".md") or os.path.exists(stem + ".docx"):
        raise SystemExit("refusing to overwrite " + stem)
    text = ""
    suffix = os.environ.get("PARTS_SUFFIX", "")  # e.g. "_joint" -> part1_joint.md ...
    for name in [f"part{i}{suffix}.md" for i in (1, 2, 3)]:
        p = os.path.join(HERE, name)
        if os.path.exists(p):
            text += open(p).read() + "\n"
        else:
            print("missing part:", name)
    slides = parse(text)
    # merged markdown with renumbered slides
    out_md = []
    for i, s in enumerate(slides, 1):
        out_md.append(f"### Slide {i} · {s['section']} · {s['kind']}")
        for f in FIELDS:
            if s.get(f):
                out_md.append(f"**{f}:** {s[f]}" if "\n" not in s[f] else f"**{f}:**\n{s[f]}")
        out_md.append("")
    open(stem + ".md", "w").write("# Deck script\n\n" + "\n".join(out_md))
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    doc.add_heading("Which story character does the Assistant take after? — deck script", 0)
    p = doc.add_paragraph(); add_runs(p, f"{len(slides)} slides. One slide per page: title, bullets, visual, footnote; speaker notes in grey.", 10, "777777")
    missing = []
    for i, s in enumerate(slides, 1):
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        lab = doc.add_paragraph(); add_runs(lab, f"Slide {i} · {s['section']} · {s['kind']}", 9, "888888")
        h = doc.add_heading(level=1); add_runs(h, s.get("Title", "(no title)"))
        for b in (s.get("Bullets") or "").splitlines():
            b = b.strip()
            if b.startswith("- "):
                par = doc.add_paragraph(style="List Bullet"); add_runs(par, b[2:], 12)
            elif b:
                par = doc.add_paragraph(); add_runs(par, b, 12)
        visual(doc, s.get("Visual", ""), missing)
        if s.get("Footnote"):
            par = doc.add_paragraph(); add_runs(par, s["Footnote"], 8.5, "777777")
        if s.get("Speaker notes"):
            par = doc.add_paragraph(); shade(par, "FFF8E6"); add_runs(par, "Speaker notes: " + s["Speaker notes"], 8.5, "555555")
    doc.save(stem + ".docx")
    print(f"wrote {stem}.md and {stem}.docx: {len(slides)} slides; missing figures: {len(missing)}")
    for m in missing:
        print("  missing:", m)


if __name__ == "__main__":
    main()
